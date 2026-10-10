"""Export helpers: CSV, JSON, XLSX, BibTeX, RIS.

Every export preserves the filters that produced it and carries a provenance
header, so a downloaded file can be traced back to the query that generated it.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from typing import Any

from openpyxl import Workbook

PROVENANCE = "provenance: IPCC Literature Database · generated {ts} · filters: {filters}"


def _now() -> str:
    return dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M:%SZ")


def _provenance_text(filters: str, metadata: dict | None = None) -> str:
    text = PROVENANCE.format(ts=_now(), filters=filters or "none")
    if metadata:
        text += " · " + " · ".join(f"{key}: {str(value).lower() if isinstance(value, bool) else value}"
                                   for key, value in metadata.items())
    return " ".join(text.split())


def to_csv(rows: list[dict], columns: list[str] | None = None,
           filters: str = "", metadata: dict | None = None) -> bytes:
    cols = columns or (list(rows[0].keys()) if rows else [])
    buf = io.StringIO()
    buf.write(f"# {_provenance_text(filters, metadata)}\n")
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({c: _cell(r.get(c)) for c in cols})
    return buf.getvalue().encode("utf-8-sig")


def to_json(rows: list[dict], columns: list[str] | None = None,
            filters: str = "", metadata: dict | None = None) -> bytes:
    payload = {"provenance": {"generated": _now(), "filters": filters or None,
                              "row_count": len(rows), **(metadata or {})},
               "columns": columns or (list(rows[0].keys()) if rows else []),
               "rows": rows}
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def to_xlsx(rows: list[dict], columns: list[str] | None = None,
            filters: str = "", sheet: str = "Export", metadata: dict | None = None) -> bytes:
    cols = columns or (list(rows[0].keys()) if rows else [])
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet[:31])
    ws.append(cols)
    for r in rows:
        ws.append([_cell(r.get(c)) for c in cols])
    meta = wb.create_sheet("Provenance")
    meta.append(["generated", _now()])
    meta.append(["filters", filters or "none"])
    meta.append(["rows", len(rows)])
    for key, value in (metadata or {}).items():
        meta.append([key, value])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _bib_key(row: dict, i: int) -> str:
    au = (row.get("authors_best") or row.get("Authors") or "")
    fam = re.split(r"[,;\s]", au.strip())[0] if au.strip() else "anon"
    fam = re.sub(r"[^A-Za-z]", "", fam) or "anon"
    yr = row.get("year_best") or row.get("Year") or "n.d."
    return f"{fam}{yr}{i}"


def to_bibtex(rows: list[dict], *, filters: str = "", metadata: dict | None = None) -> bytes:
    parts = ["% " + _provenance_text(filters, metadata)]
    for i, r in enumerate(rows, 1):
        doi = r.get("doi_norm") or r.get("DOI (Extracted)") or ""
        doi = str(doi).replace("https://doi.org/", "").strip()
        entry = "article" if (r.get("journal_best") or r.get("Publisher")) else "misc"
        fields = {
            "title": r.get("title_best") or r.get("Article Title") or "",
            "author": (r.get("authors_best") or r.get("Authors") or "").replace(";", " and"),
            "year": r.get("year_best") or r.get("Year") or "",
            "journal": r.get("journal_best") or r.get("Publisher") or "",
            "doi": doi,
            "note": f"citedInIPCC: {r.get('report_id') or r.get('Report') or ''} "
                    f"{r.get('wg_label') or r.get('Working Group') or ''} "
                    f"{r.get('chapter_label') or r.get('Chapter') or ''}".strip(),
        }
        body = ",\n".join(f"  {k} = {{{v}}}" for k, v in fields.items() if v)
        parts.append(f"@{entry}{{{_bib_key(r, i)},\n{body}\n}}")
    return ("\n\n".join(parts) + "\n").encode("utf-8")


def to_ris(rows: list[dict], *, filters: str = "", metadata: dict | None = None) -> bytes:
    out = []
    provenance = _provenance_text(filters, metadata)
    for r in rows:
        doi = str(r.get("doi_norm") or r.get("DOI (Extracted)") or "").replace(
            "https://doi.org/", "").strip()
        lines = ["TY  - JOUR" if (r.get("journal_best") or r.get("Publisher")) else "TY  - GEN"]
        title = r.get("title_best") or r.get("Article Title")
        if title:
            lines.append(f"TI  - {title}")
        authors = r.get("authors_best") or r.get("Authors") or ""
        for a in [x.strip() for x in re.split(r";", authors) if x.strip()]:
            lines.append(f"AU  - {a}")
        if r.get("year_best") or r.get("Year"):
            lines.append(f"PY  - {r.get('year_best') or r.get('Year')}")
        if r.get("journal_best") or r.get("Publisher"):
            lines.append(f"JO  - {r.get('journal_best') or r.get('Publisher')}")
        if doi:
            lines.append(f"DO  - {doi}")
            lines.append(f"UR  - https://doi.org/{doi}")
        ctx = " ".join(str(r.get(k) or "") for k in
                       ("report_id", "Report", "wg_label", "Working Group",
                        "chapter_label", "Chapter")).strip()
        if ctx:
            lines.append(f"N1  - cited in IPCC: {ctx}")
        lines.append(f"N1  - {provenance}")
        lines.append("ER  - ")
        out.append("\n".join(lines))
    return ("\n\n".join(out) + "\n").encode("utf-8")


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", str(v))[:32000]
